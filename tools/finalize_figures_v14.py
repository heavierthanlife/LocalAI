# -*- coding: utf-8 -*-
"""基于作者手工修改的 v1.2 docx：①替换 16 张图为 v14 重制版；②补齐/更新图注（图内注已删）。"""
import os, sys, hashlib, glob, copy
import docx
from docx.shared import Emu
from docx.image.image import Image
from docx.text.paragraph import Paragraph

SRC = sys.argv[1]
DST = sys.argv[2]
OLD_DIR = "build/figures_final"
NEW_DIR = "build/figures_v14"

def md5(b): return hashlib.md5(b).hexdigest()
oldmap = {md5(open(f,'rb').read()): os.path.basename(f) for f in glob.glob(OLD_DIR+'/*.png')}
newmap = {os.path.basename(f): f for f in glob.glob(NEW_DIR+'/*.png')}

doc = docx.Document(SRC)

# ---------- 1) 替换图片 ----------
replaced = []
for i, sh in enumerate(doc.inline_shapes):
    rId = sh._inline.graphic.graphicData.pic.blipFill.blip.embed
    part = doc.part.related_parts[rId]
    name = oldmap.get(md5(part.blob))
    if not name or name not in newmap:
        print("SKIP", i, name); continue
    data = open(newmap[name], 'rb').read()
    img = Image.from_blob(data)
    part._blob = data
    w = sh.width
    if w.cm > 15.0: w = Emu(int(15.0*360000))
    hgt = int(w * img.px_height / img.px_width)
    maxh = int(21.5*360000)
    if hgt > maxh:
        hgt = maxh; w = Emu(int(hgt * img.px_width / img.px_height))
    sh.width, sh.height = w, Emu(hgt)
    replaced.append(name)
print("images replaced:", len(replaced))

# ---------- 2) 图注：更新/插入 ----------
paras = doc.paragraphs
def find_par(pred):
    for p in paras:
        if pred(p.text.strip()): return p
    return None

def set_text(p, text):
    runs = p.runs
    if not runs:
        r = p.add_run(); r.text = text
    else:
        runs[0].text = text
        for r in runs[1:]: r.text = ""
    return p

# 2a 更新已有图注
updates = {
 "注：横读为同一维度在诊断、归因、对策与评价中的对应关系。":
 "注：横读为同一维度在诊断、归因、对策与评价中的对应关系；纵读为同一维度从现状到指标的逻辑链条。",
 "注：权重在后续实测阶段以熵权法与AHP组合赋权确定；本文情景测算采用等权标准化合成。":
 "注：权重在后续实测阶段以熵权法与AHP组合赋权确定，本文情景测算采用等权标准化合成；评价方法为情景测算与敏感性分析（后续实测阶段拟采用配对检验与中断时间序列）。",
 "注：数据经聚合与四舍五入处理；清标耗时为对数刻度。":
 "注：数据经公司许可、聚合与四舍五入处理；清标前后计时口径不完全一致，仅作方向性比较；清标耗时为对数刻度。",
}
for old,new in updates.items():
    p = find_par(lambda t, o=old: t == o)
    if p: set_text(p,new); print("updated note:", new[:20])
    else: print("WARN note not found:", old[:20])

# 2b 为图1-1、图4-1 插入图注（模板取图2-1的注）
tpl = find_par(lambda t: t.startswith("注：据第2章理论阐述"))
if tpl is None: tpl = find_par(lambda t: t.startswith("注："))
def insert_note_after(caption_text, note_text):
    cap = find_par(lambda t, c=caption_text: t == c)
    if not cap:
        print("WARN caption not found:", caption_text); return
    # 已有注则更新
    nxt = cap._p.getnext()
    if nxt is not None:
        np = Paragraph(nxt, cap._parent)
        if np.text.strip().startswith("注："):
            set_text(np, note_text); print("insert->update", caption_text); return
    new_el = copy.deepcopy(tpl._p)
    cap._p.addnext(new_el)
    np = Paragraph(new_el, cap._parent)
    set_text(np, note_text)
    print("inserted note after:", caption_text)

insert_note_after("图1-1 研究技术路线图", "注：据研究内容、研究方法与技术路线整理绘制。")
insert_note_after("图4-1 工作流程问题成因鱼骨图", "注：据问卷调查与半结构化访谈结果归纳绘制。")

doc.save(DST)
print("saved", DST)
