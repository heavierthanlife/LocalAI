import docx, sys
p = sys.argv[1]
d = docx.Document(p)
for i, para in enumerate(d.paragraphs):
    t = para.text.strip()
    if t:
        print(f"[{i}][{para.style.name}] {t}")
print(f"\n=== PARAS {len(d.paragraphs)} TABLES {len(d.tables)} ===")
for ti, tbl in enumerate(d.tables):
    print(f"--- table {ti}: {len(tbl.rows)}x{len(tbl.columns)} ---")
    for r in tbl.rows[:3]:
        print(" | ".join(c.text.strip() for c in r.cells))
