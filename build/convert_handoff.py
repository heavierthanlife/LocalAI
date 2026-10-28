# -*- coding: utf-8 -*-
"""把 HANDOFF 的 markdown 转成 docx（简易转换器：标题/表格/代码块/列表/引用）。"""
import re
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

SRC = "00_HANDOFF_交接说明.md"
DST = "00_HANDOFF_交接说明.docx"
SONG, HEI, MONO = "宋体", "黑体", "Consolas"

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.left_margin = sec.right_margin = Cm(2.4)
sec.top_margin = sec.bottom_margin = Cm(2.4)

st = doc.styles["Normal"]
st.font.name = SONG
st.font.size = Pt(10.5)
st.element.rPr.rFonts.set(qn("w:eastAsia"), SONG)
st.paragraph_format.line_spacing = 1.35
st.paragraph_format.space_after = Pt(2)

for name, size in [("Heading 1", 16), ("Heading 2", 13.5), ("Heading 3", 12)]:
    s = doc.styles[name]
    s.font.name = HEI
    s.font.size = Pt(size)
    s.font.bold = True
    s.font.color.rgb = RGBColor(0x1F, 0x38, 0x64)
    s.element.rPr.rFonts.set(qn("w:eastAsia"), HEI)
    s.paragraph_format.space_before = Pt(10)
    s.paragraph_format.space_after = Pt(4)


def inline(text):
    text = re.sub(r"\*\*(.+?)\*\*", r"\1", text)
    text = re.sub(r"`(.+?)`", r"\1", text)
    text = re.sub(r"\[(.+?)\]\((.+?)\)", r"\1", text)
    return text


def add_par(text, style="Normal", size=10.5, bold=False, indent=0, align=WD_ALIGN_PARAGRAPH.JUSTIFY,
            name=SONG, color=None, space_after=2):
    p = doc.add_paragraph(style=style)
    pf = p.paragraph_format
    pf.first_line_indent = Pt(0)
    pf.left_indent = Pt(indent)
    pf.space_after = Pt(space_after)
    pf.alignment = align
    if text:
        r = p.add_run(text)
        r.font.name = name
        r.font.size = Pt(size)
        r.bold = bold
        r._element.rPr.rFonts.set(qn("w:eastAsia"), name)
        if color:
            r.font.color.rgb = color
    return p


def cell_border(cell, sz=6):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "bottom"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single"); el.set(qn("w:sz"), str(sz)); el.set(qn("w:color"), "000000")
        borders.append(el)
    tcPr.append(borders)


lines = open(SRC, encoding="utf-8").read().split("\n")
i = 0
first_h1 = True
while i < len(lines):
    line = lines[i]
    s = line.strip()

    # 代码块
    if s.startswith("```"):
        i += 1
        buf = []
        while i < len(lines) and not lines[i].strip().startswith("```"):
            buf.append(lines[i])
            i += 1
        i += 1
        for code in buf:
            p = add_par(code if code.strip() else " ", size=9, name=MONO,
                        align=WD_ALIGN_PARAGRAPH.LEFT, space_after=0)
            pf = p.paragraph_format
            pf.left_indent = Pt(10)
            shd = OxmlElement("w:shd")
            shd.set(qn("w:val"), "clear"); shd.set(qn("w:fill"), "F2F2F2")
            p._p.get_or_add_pPr().append(shd)
        add_par(" ", size=6)
        continue

    # 表格
    if s.startswith("|") and i + 1 < len(lines) and set(lines[i + 1].replace("|", "").replace("-", "").replace(":", "").strip()) <= {" "}:
        rows = []
        while i < len(lines) and lines[i].strip().startswith("|"):
            rows.append([c.strip() for c in lines[i].strip().strip("|").split("|")])
            i += 1
        rows = [r for r in rows if not set("".join(r)) <= set("-: ")]
        if rows:
            ncol = max(len(r) for r in rows)
            t = doc.add_table(rows=len(rows), cols=ncol)
            t.alignment = WD_TABLE_ALIGNMENT.CENTER
            tblPr = t._tbl.tblPr
            borders = OxmlElement("w:tblBorders")
            for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
                el = OxmlElement(f"w:{edge}"); el.set(qn("w:val"), "none"); borders.append(el)
            tblPr.append(borders)
            for ri, row in enumerate(rows):
                for ci in range(ncol):
                    cell = t.cell(ri, ci)
                    cell.text = ""
                    pp = cell.paragraphs[0]
                    pp.paragraph_format.first_line_indent = Pt(0)
                    pp.paragraph_format.space_after = Pt(1)
                    pp.paragraph_format.space_before = Pt(1)
                    pp.paragraph_format.line_spacing = 1.0
                    val = inline(row[ci]) if ci < len(row) else ""
                    r = pp.add_run(val)
                    r.font.name = SONG; r.font.size = Pt(9); r.bold = (ri == 0)
                    r._element.rPr.rFonts.set(qn("w:eastAsia"), SONG)
                    cell_border(cell, 8 if ri == 0 else (8 if ri == len(rows) - 1 else 4))
            add_par(" ", size=6)
        continue

    if not s:
        i += 1
        continue
    if s.startswith("---"):
        i += 1
        continue

    # 标题
    m = re.match(r"^(#{1,3})\s+(.*)$", s)
    if m:
        lvl, text = len(m.group(1)), inline(m.group(2))
        if lvl == 1 and first_h1:
            p = doc.add_paragraph(style="Heading 1")
            p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run(text)
            first_h1 = False
        else:
            doc.add_paragraph(text, style=f"Heading {lvl}")
        i += 1
        continue

    # 引用
    if s.startswith(">"):
        add_par(inline(s.lstrip("> ")), size=9.5, indent=12, color=RGBColor(0x59, 0x59, 0x59))
        i += 1
        continue

    # 列表
    m = re.match(r"^(\s*)[-*]\s+(.*)$", line)
    if m:
        depth = len(m.group(1)) // 2
        add_par("· " + inline(m.group(2)), indent=12 + 12 * depth)
        i += 1
        continue

    add_par(inline(s))
    i += 1

doc.save(DST)
print("saved", DST)
