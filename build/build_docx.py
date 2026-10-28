# -*- coding: utf-8 -*-
"""基于原论文生成 v0.5 修订稿：插入图1–图11（章序编号）+ 修正标题层级 + 修正明显笔误。
原文件 thesis_orig.docx 保持不动，输出 02_..._v0.5_含图表.docx
"""
import copy, os
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

SRC = "thesis_orig.docx"
DST = "02_姚烨_2024245834_Z公司工作流程数字化再造研究_v0.5_含图表.docx"
FIG = "figures"

doc = Document(SRC)
paras = doc.paragraphs


def find(prefix, start=0):
    for p in paras[start:]:
        if p.text.strip().startswith(prefix):
            return p
    raise KeyError(prefix)


# ---------------- 1. 文字修订 ----------------
TEXT_FIXES = [
    ("2021年，中国提出“双碳”目标", "2020年9月，中国提出“双碳”目标"),
    ("本研究Z公司为案例", "本研究以Z公司为案例"),
    ("1.3..3行业实证分析与可持续转型启示研究", "1.3.3 行业实证分析与可持续转型启示研究"),
    ("首先，构建数字化碳监测平台，通过模块开发、物联网集成与试点测试，实现碳足迹全生命周期追踪，提升监测精度与服务能力。",
     "首先，构建数字化工作活动平台，通过流程引擎与模块集成、试点测试，实现工作流程全周期追踪，提升流程效率与服务能力。"),
]
for old, new in TEXT_FIXES:
    hit = 0
    for p in doc.paragraphs:
        if old in p.text:
            # 尽量在 run 级别替换，保留原有格式
            for r in p.runs:
                if old in r.text:
                    r.text = r.text.replace(old, new)
                    hit += 1
            if hit == 0 and old in p.text:
                p.text = p.text.replace(old, new)
                hit += 1
    print(f"[fix] {old[:24]}… -> {'OK' if hit else 'MISS'}")

# ---------------- 2. 标题层级 ----------------
H = {
    "2.3本章小结": 2, "3.2 存在问题": 2,
    "3.2.1 数字化工具集成度低": 3, "3.2.2 流程冗余导致资源浪费": 3,
    "3.2.3 员工数字化意识与技能不足": 3, "3.3本章小结": 2,
    "1.3.1": 3, "1.3.2": 3, "1.3.3": 3, "1.3.4": 3, "1.4本章小结": 2,
    "4.5本章小结": 2, "5.5本章小结": 2,
    "第6章 结论": 1, "6.1 研究结论": 2, "6.2 研究展望": 2,
    "参考文献": 1, "附录 A": 1, "附录 B": 1,
}
for p in doc.paragraphs:
    t = p.text.strip()
    for key, lvl in H.items():
        if t.startswith(key):
            p.style = doc.styles[f"Heading {lvl}"]
            for r in p.runs:
                r.bold = True
                r.font.color.rgb = RGBColor(0, 0, 0)
                r.font.size = None
            print(f"[heading] H{lvl} <- {t[:30]}")
            break

# ---------------- 3. 插入图表 ----------------
def _mk_para_after(anchor_p, kind, **kw):
    """在 anchor 之后插入新段落；返回新段落。kind: 'img' / 'cap' / 'note'"""
    p = doc.add_paragraph()
    if kind == "img":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pf = p.paragraph_format
        pf.space_before = Pt(6)
        pf.space_after = Pt(2)
        pf.line_spacing = 1.0
        run = p.add_run()
        run.add_picture(kw["path"], width=Cm(kw["width"]))
    else:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pf = p.paragraph_format
        pf.space_before = Pt(0)
        pf.space_after = Pt(6 if kind == "note" else 3)
        pf.line_spacing = 1.0
        run = p.add_run(kw["text"])
        run.font.name = "宋体"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
        if kind == "cap":
            run.font.size = Pt(10.5)
            run.bold = True
        else:
            run.font.size = Pt(8.5)
            run.font.color.rgb = RGBColor(0x59, 0x59, 0x59)
    anchor_p._p.addnext(p._p)
    return p


FIGURES = [
    # (锚点段落前缀, 图片, 图题, 注释, 宽度cm)
    ("研究方法的选择注重定性与定量的结合", "图01_研究技术路线图.png",
     "图1-1 研究技术路线图", None, 11.0),
    ("这些理论相互补充：BPR提供具体再造方法", "图02_理论整合分析框架.png",
     "图2-1 业务流程再造—TOE—动态能力整合分析框架",
     "注：据第2章理论阐述与第3—5章分析逻辑整理绘制。", 14.5),
    ("这些理论相互补充：BPR提供具体再造方法", "图03_TOE维度映射矩阵.png",
     "图2-2 TOE维度与问题、成因、对策及观测指标的映射矩阵",
     "注：横读为同一维度在诊断、归因、对策与评价中的对应关系。", 15.0),
    ("Z公司的工作流程主要围绕资产评估业务展开", "图04_AS-IS六阶段评估流程.png",
     "图3-1 AS-IS六阶段评估流程与瓶颈环节",
     "注：据3.1.2节流程描述绘制，问题标注来自问卷调查结果。", 15.0),
    ("总体样本大小适中，抽样误差控制在5%以内", "图07_调查样本结构.png",
     "图3-2 调查样本结构（N=185）",
     "数据来源：问卷调查结果，对应表3-1至表3-7。", 15.0),
    ("该表显示云计算应用最广泛，但集成度最低", "图08_数字化工具使用率与集成度.png",
     "图3-3 数字化工具使用率、满意度均值与离散度",
     "数据来源：表3-8。", 14.0),
    ("综上，调查结果揭示Z公司工作流程现状的二元性", "图05_核心问题量化表征.png",
     "图3-4 三大核心问题的量化表征",
     "数据来源：3.1.2节与3.2节调查结果；资源浪费率为访谈估算值。", 14.0),
    ("公司作为行业头部机构，员工整体素质高", "图06_问题成因鱼骨图.png",
     "图4-1 工作流程问题成因鱼骨图",
     "注：据第4章四类成因归纳绘制。", 15.0),
    ("上述成因并非孤立存在，彼此之间呈现相互交织", "图09_问题-成因-对策传导链.png",
     "图4-2 问题—成因—对策的传导链条",
     "注：据第3—5章内容整理；灰线表示成因对问题的归因关系。", 15.0),
    ("总体而言，这一对策通过再造与集成，实现工作流程的数字化优化", "图10_再造前后目标指标对比.png",
     "图5-1 再造前后三项关键指标的目标对比（AS-IS=100）",
     "注：目标值据第5章方案测算——流程周期降15%—20%、复核由3轮精简为单级、纸质文档降至10%以下；均非实测值，正式稿须以系统上线后实测数据替换。", 14.0),
    ("最后，建立培训体系的持续优化机制", "图11_再造效果评价指标体系.png",
     "图5-2 再造效果评价指标体系框架",
     "注：指标权重待以熵权法与AHP组合赋权测算。", 15.0),
]

prev_anchor = None
last_tail = None
for anchor, img, cap, note, width in FIGURES:
    if anchor == prev_anchor and last_tail is not None:
        a = last_tail
    else:
        a = find(anchor)
    path = os.path.join(FIG, img)
    assert os.path.exists(path), path
    # 顺序：img 插在锚点后；cap 插在 img 后；note 插在 cap 后
    _mk_para_after(a, "img", path=path, width=width)
    # 重新取刚插入的图片段落需通过 XML 顺序：锚点后第一个 w:p
    img_p = a._p.getnext()
    from docx.text.paragraph import Paragraph
    img_p_obj = Paragraph(img_p, a._parent)
    _mk_para_after(img_p_obj, "cap", text=cap)
    cap_p = img_p.getnext()
    last_tail = Paragraph(cap_p, a._parent)
    if note:
        last_tail = _mk_para_after(last_tail, "note", text=note)
    prev_anchor = anchor
    print(f"[fig] {cap}")

doc.save(DST)
print("saved", DST)
