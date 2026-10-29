# -*- coding: utf-8 -*-
"""组装 v1.0 论文 docx：封面 + 正文 + 表 + 图 + 参考文献（沿用原 63 篇）+ 附录"""
import os
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

from text_part1 import BLOCKS_1
from text_part2 import BLOCKS_2, BLOCKS_2B, BLOCKS_3
from text_part3 import BLOCKS_4, BLOCKS_5, BLOCKS_6, BLOCKS_APPENDIX

SRC = "thesis_orig.docx"
DST = "07_姚烨_2024245834_Z公司工作流程数字化再造研究_v1.0_招标代理案例版.docx"
FIG = "figures"

SONG, HEI = "宋体", "黑体"

# ------------------------------------------------------------------ 表格数据
TABLES = {
    "t31": ("表3-1 调查样本性别分布", [["性别", "人数", "占比(%)"],
                                     ["男性", "105", "56.8"], ["女性", "80", "43.2"]]),
    "t32": ("表3-2 调查样本年龄分布", [["年龄段", "人数", "占比(%)"],
                                     ["30岁以下", "53", "28.6"], ["31-40岁", "84", "45.4"],
                                     ["41-50岁", "38", "20.5"], ["51岁以上", "10", "5.4"]]),
    "t33": ("表3-3 调查样本教育背景分布", [["教育背景", "人数", "占比(%)"],
                                        ["大专及以下", "14", "7.6"], ["本科", "106", "57.3"],
                                        ["硕士及以上", "65", "35.1"]]),
    "t34": ("表3-4 调查样本工作年限分布", [["工作年限", "人数", "占比(%)"],
                                        ["5年以下", "60", "32.4"], ["6-10年", "76", "41.1"],
                                        ["11年以上", "49", "26.5"]]),
    "t35": ("表3-5 调查样本职位分布", [["职位级别", "人数", "占比(%)"],
                                    ["普通员工", "109", "58.9"], ["中层管理人员", "51", "27.6"],
                                    ["高层管理人员", "25", "13.5"]]),
    "t36": ("表3-6 调查样本部门分布", [["部门", "人数", "占比(%)"],
                                    ["招标代理部", "79", "42.7"], ["造价咨询部", "47", "25.4"],
                                    ["管理部", "35", "18.9"], ["其他部门", "24", "13.0"]]),
    "t37": ("表3-7 调查样本地域分布", [["地域", "人数", "占比(%)"],
                                    ["总部（北京）", "121", "65.4"], ["分支机构", "64", "34.6"]]),
    "t38": ("表3-8 数字化工具应用分布", [["工具类型", "使用率(%)", "均值", "标准差"],
                                       ["云计算/大数据", "72.4", "4.1", "0.9"],
                                       ["人工智能/自动化", "58.9", "3.6", "1.2"],
                                       ["区块链/安全工具", "45.4", "3.1", "1.3"],
                                       ["整体集成度", "45.9", "3.2", "1.1"]]),
    "t39": ("表3-9 流程与资源效率关键指标", [["指标", "均值", "标准差", "问题占比(%)"],
                                          ["纸质文档使用", "3.4", "1.1", "40.0"],
                                          ["节点延时", "3.2", "1.2", "52.4"],
                                          ["整体效率满意度", "3.5", "1.0", "—"]]),
    "t310": ("表3-10 人力与激励状况关键指标", [["指标", "均值", "标准差", "问题占比(%)"],
                                            ["数字化意识", "3.7", "1.0", "—"],
                                            ["培训覆盖率", "3.1", "1.3", "55.0"],
                                            ["激励机制认同", "3.1", "1.1", "40.5"]]),
    "t51": ("表5-1 平台主要技术选型", [["层次", "选型", "说明"],
                                    ["后端框架", "Flask 3.1 + 应用工厂", "按业务域拆分蓝图与服务层"],
                                    ["数据存储", "PostgreSQL 16", "76 张表，幂等建表 + 版本化迁移"],
                                    ["缓存与消息", "Redis 7 + Celery 5", "异步任务、定时调度与缓存"],
                                    ["智能编排", "LangGraph（状态图代理）", "多步推理与流式输出"],
                                    ["模型服务", "双供应商大模型接入", "私有化部署，数据不出内网"],
                                    ["检索增强", "ChromaDB + 向量模型", "知识库语义检索与溯源"],
                                    ["文档处理", "PyMuPDF / python-docx / OCR", "多格式解析与扫描件识别"],
                                    ["部署运维", "Docker 编排 + Nginx", "六服务容器化，HTTPS 接入"]]),
    "t52": ("表5-2 关键功能模块与实现要点", [["模块", "核心能力", "主要输出"],
                                          ["文档智能解析", "多格式解析、章节条款切分、OCR 识别", "结构化文本、条款索引"],
                                          ["清标辅助评分", "五维度要素提取、加权合成评分", "评分表、问题清单"],
                                          ["围串标风险识别", "文本相似度、报价异常、关联网络", "风险提示清单"],
                                          ["合规检查", "法规条款结构化、增量检索与校验", "合规检查意见"],
                                          ["智能对话与知识库", "检索增强生成、引用溯源", "问答结果、条款引用"],
                                          ["项目协同与信用查询", "文件版本、评论待办、信用采集", "项目档案、信用报告"]]),
    "t61": ("表6-1 再造效果评价指标与测量口径", [["准则层", "指标", "测量口径", "数据来源"],
                                             ["流程效率", "流程周期", "实际结束日期 − 实际开始日期（天）", "平台时间轴"],
                                             ["流程效率", "计划偏差率", "(实际周期 − 计划周期) / 计划周期", "平台时间轴"],
                                             ["流程效率", "里程碑延误", "实际日期 − 计划日期（天）", "里程碑记录"],
                                             ["流程效率", "返工强度", "修订轮次、生成次数", "流程 KPI 表"],
                                             ["流程效率", "任务响应", "待办完成时间 − 创建时间（小时）", "待办记录"],
                                             ["集成水平", "数据一次录入率", "无需重复录入的字段 / 总字段", "系统日志"],
                                             ["集成水平", "自动流转占比", "接口自动流转单据 / 总单据", "接口日志"],
                                             ["集成水平", "接口成功率", "成功调用 / 总调用", "接口日志"],
                                             ["资源节约", "单项目人工工时", "项目工时合计 / 项目数", "业务台账"],
                                             ["资源节约", "纸质文档占比", "纸质流转文件 / 总文件", "档案记录"],
                                             ["资源节约", "人工搬运次数", "每项目跨系统手工导入导出次数", "问卷 + 访谈"],
                                             ["业务价值", "节点准时交付率", "按期完成节点 / 总节点", "里程碑记录"],
                                             ["业务价值", "客户满意度", "满意度评分均值", "客户回访"],
                                             ["业务价值", "单项目毛利率", "(收入 − 成本) / 收入", "财务台账"],
                                             ["组织能力", "培训覆盖率", "受训人数 / 应训人数", "培训记录"],
                                             ["组织能力", "工具采纳率", "月活跃用户 / 应使用用户", "系统日志"],
                                             ["组织能力", "留痕完整率", "有审计记录的关键操作 / 关键操作总数", "审计日志"]]),
    "t62": ("表6-2 情景参数设定（相对基线的变化率）", [["指标", "保守情景", "基准情景", "乐观情景"],
                                                ["流程周期", "-8%", "-15%", "-22%"],
                                                ["清标环节耗时", "-18%", "-32%", "-45%"],
                                                ["文件返工轮次", "-20%", "-40%", "-55%"],
                                                ["人工数据搬运次数", "-25%", "-45%", "-60%"]]),
}

FIG_W = {}
TOTAL = 0


def set_font(run, name=SONG, size=12, bold=False, color=None):
    run.font.name = name
    run.font.size = Pt(size)
    run.bold = bold
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if color:
        run.font.color.rgb = color


def style_setup(doc):
    st = doc.styles["Normal"]
    st.font.name = SONG
    st.font.size = Pt(12)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), SONG)
    pf = st.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    pf.space_before, pf.space_after = Pt(0), Pt(0)
    pf.first_line_indent = Pt(24)
    for name, size, align in [("Heading 1", 16, WD_ALIGN_PARAGRAPH.CENTER),
                              ("Heading 2", 14, WD_ALIGN_PARAGRAPH.LEFT),
                              ("Heading 3", 12, WD_ALIGN_PARAGRAPH.LEFT),
                              ("Heading 4", 12, WD_ALIGN_PARAGRAPH.LEFT)]:
        s = doc.styles[name]
        s.font.name = HEI
        s.font.size = Pt(size)
        s.font.bold = True
        s.font.color.rgb = RGBColor(0, 0, 0)
        s.element.rPr.rFonts.set(qn("w:eastAsia"), HEI)
        s.paragraph_format.first_line_indent = Pt(0)
        s.paragraph_format.space_before = Pt(12 if name == "Heading 1" else 8)
        s.paragraph_format.space_after = Pt(8 if name == "Heading 1" else 4)
        s.paragraph_format.alignment = align


def add_par(doc, text, style="Normal", size=12, bold=False, indent=True,
            align=WD_ALIGN_PARAGRAPH.JUSTIFY, name=SONG, space_after=0, line=1.5):
    p = doc.add_paragraph(style=style)
    pf = p.paragraph_format
    pf.first_line_indent = Pt(size * 2) if indent else Pt(0)
    pf.alignment = align
    pf.line_spacing = line
    pf.space_after = Pt(space_after)
    if text:
        set_font(p.add_run(text), name, size, bold)
    return p


def set_cell_borders(cell, top=None, bottom=None):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = tcPr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tcPr.append(borders)
    for edge, sz in (("top", top), ("bottom", bottom)):
        if sz is None:
            continue
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(sz))
        el.set(qn("w:color"), "000000")
        borders.append(el)


def add_table(doc, key):
    title, data = TABLES[key]
    # 表题
    p = doc.add_paragraph()
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Pt(0)
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    set_font(p.add_run(title), SONG, 10.5, True)

    nrow, ncol = len(data), len(data[0])
    t = doc.add_table(rows=nrow, cols=ncol)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = True
    # 去掉默认样式边框
    tbl = t._tbl
    tblPr = tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "none")
        borders.append(el)
    tblPr.append(borders)
    # 表头跨页重复
    trPr = t.rows[0]._tr.get_or_add_trPr()
    th = OxmlElement("w:tblHeader"); th.set(qn("w:val"), "true"); trPr.append(th)
    for i, row in enumerate(data):
        for j, val in enumerate(row):
            cell = t.cell(i, j)
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.first_line_indent = Pt(0)
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.0
            set_font(p.add_run(val), SONG, 10.5, i == 0)
            top = 12 if i == 0 else None
            bottom = 6 if i == 0 else (12 if i == nrow - 1 else None)
            set_cell_borders(cell, top=top, bottom=bottom)
    # 表后空行
    q = doc.add_paragraph()
    q.paragraph_format.first_line_indent = Pt(0)
    q.paragraph_format.space_after = Pt(0)
    return t


def add_fig(doc, fname, caption, note, width):
    path = os.path.join(FIG, fname)
    assert os.path.exists(path), path
    p = doc.add_paragraph()
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Pt(0)
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.0
    p.add_run().add_picture(path, width=Cm(width))
    c = doc.add_paragraph()
    c.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    c.paragraph_format.first_line_indent = Pt(0)
    c.paragraph_format.space_after = Pt(8 if note else 10)
    c.paragraph_format.line_spacing = 1.0
    set_font(c.add_run(caption), SONG, 10.5, True)
    if note:
        n = doc.add_paragraph()
        n.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        n.paragraph_format.first_line_indent = Pt(0)
        n.paragraph_format.space_after = Pt(10)
        n.paragraph_format.line_spacing = 1.0
        set_font(n.add_run(note), SONG, 8.5, False, RGBColor(0x59, 0x59, 0x59))


def add_page_number(section):
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    run._r.addnext(fld)


def build():
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(3.0)
    sec.top_margin = sec.bottom_margin = Cm(3.0)
    style_setup(doc)
    add_page_number(sec)

    # ---------------- 封面 ----------------
    p = doc.add_paragraph(); p.paragraph_format.first_line_indent = Pt(0)
    set_font(p.add_run("中图分类号：　　　　　　　　　　　　　　　"), SONG, 10.5)
    doc.add_paragraph()
    t = doc.add_table(rows=7, cols=2)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    cover = [("论文类型", "专题研究类论文"), ("题目", "Z公司工作流程数字化再造研究"),
             ("专业或领域", "工商管理"), ("研究方向", ""),
             ("专业硕士生", "姚　烨　2024245834"), ("指导教师", ""), ("现场导师", "")]
    for i, (k, v) in enumerate(cover):
        for j, val in enumerate((k, v)):
            cell = t.cell(i, j)
            cell.text = ""
            pp = cell.paragraphs[0]
            pp.paragraph_format.first_line_indent = Pt(0)
            pp.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER if j == 0 else WD_ALIGN_PARAGRAPH.LEFT
            pp.paragraph_format.line_spacing = 1.5
            set_font(pp.add_run(val), SONG, 12, j == 0)
            set_cell_borders(cell, top=6, bottom=6)
    for _ in range(4):
        doc.add_paragraph()
    p = doc.add_paragraph(); p.paragraph_format.first_line_indent = Pt(0)
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(p.add_run("二○二六年十月"), SONG, 14, True)
    doc.add_page_break()

    # ---------------- 正文 ----------------
    for blk in BLOCKS_1 + BLOCKS_2 + BLOCKS_2B + BLOCKS_3 + BLOCKS_4 + BLOCKS_5 + BLOCKS_6:
        kind = blk[0]
        if kind == "h1":
            p = doc.add_paragraph(blk[1], style="Heading 1")
            t = blk[1].strip()
            if t == "摘    要" or t.startswith("第") or t.startswith("附录") or t in ("ABSTRACT",):
                p.paragraph_format.page_break_before = True
        elif kind == "h2":
            doc.add_paragraph(blk[1], style="Heading 2")
        elif kind == "h3":
            doc.add_paragraph(blk[1], style="Heading 3")
        elif kind == "h4":
            doc.add_paragraph(blk[1], style="Heading 4")
        elif kind == "p":
            add_par(doc, blk[1])
        elif kind == "tbl":
            add_table(doc, blk[1])
        elif kind == "tcap":
            pass  # 表题已在 add_table 内处理
        elif kind == "fig":
            add_fig(doc, blk[1], blk[2], blk[3], blk[4])
        else:
            raise ValueError(kind)

    # ---------------- 参考文献（沿用原 63 篇） ----------------
    _rp = doc.add_paragraph("参考文献", style="Heading 1")
    _rp.paragraph_format.page_break_before = True
    src = Document(SRC)
    refs = [p.text.strip() for p in src.paragraphs if p.text.strip().startswith("[") and "]" in p.text.strip()[:6]]
    for r in refs:
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.first_line_indent = Pt(0)
        pf.left_indent = Pt(21)
        pf.first_line_indent = Pt(-21)
        pf.line_spacing = 1.3
        pf.space_after = Pt(2)
        pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        set_font(p.add_run(r), SONG, 10.5)
    print("refs:", len(refs))

    # ---------------- 附录 ----------------
    for blk in BLOCKS_APPENDIX:
        kind = blk[0]
        if kind == "h1":
            doc.add_paragraph(blk[1], style="Heading 1")
        elif kind == "h2":
            doc.add_paragraph(blk[1], style="Heading 2")
        elif kind == "p":
            add_par(doc, blk[1], indent=(not blk[1].startswith(("□", "1 □", "1.", "2.", "3.", "4.", "5.",
                                                               "6.", "7.", "8.", "9.", "10.", "11.",
                                                               "12.", "13.", "14.", "15.", "16.", "17.",
                                                               "18.", "19.", "20.", "21.", "22.", "23.",
                                                               "24.", "25.", "-", "访谈", "调查", "填写",
                                                               "说明", "亲爱的", "结束语", "Z公司",
                                                               "感谢"))))

    doc.save(DST)
    print("saved", DST)


build()
