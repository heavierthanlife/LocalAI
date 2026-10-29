# -*- coding: utf-8 -*-
"""生成 MBA 学位论文答辩 PPT（16:9），内容取自终版 PDF 与论文数据。"""
import os
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from PIL import Image

NAVY = RGBColor(0x12, 0x3A, 0x63)
BLUE = RGBColor(0x2E, 0x6D, 0xA4)
GOLD = RGBColor(0xC9, 0xA2, 0x27)
GRAY = RGBColor(0x5A, 0x6B, 0x7B)
LIGHT = RGBColor(0xEA, 0xF1, 0xF8)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
DARK = RGBColor(0x1B, 0x2A, 0x38)
FONT = "微软雅黑"

FIG = "figs_style_thesis_a"
FIGS = {
 "route": "图01_研究技术路线图.png",
 "framework": "图02_理论整合分析框架.png",
 "toe": "图03_TOE维度映射矩阵.png",
 "asis": "图04_AS-IS六阶段评估流程.png",
 "sample": "图07_调查样本结构.png",
 "tools": "图08_数字化工具使用率与集成度.png",
 "problem": "图05_核心问题量化表征.png",
 "fishbone": "图06_问题成因鱼骨图.png",
 "chain": "图09_问题-成因-对策传导链.png",
 "arch": "图12_系统总体技术架构.png",
 "data": "图13_流程数据模型与算指标.png" if os.path.exists("figs_style_thesis_a/图13_流程数据模型与算指标.png") else "图13_流程数据模型与可算指标.png",
 "kpi": "图11_再造效果评价指标体系.png",
 "rework": "图17_返工轮次基线.png",
 "measured": "图18_再造前后实测对比.png",
 "scenario": "图16_情景测算结果.png",
 "target": "图10_再造前后目标指标对比.png",
}

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
SW, SH = prs.slide_width, prs.slide_height
BLANK = prs.slide_layouts[6]


def set_font(run, size, bold=False, color=DARK, name=FONT):
    f = run.font
    f.name = name; f.size = Pt(size); f.bold = bold; f.color.rgb = color
    rPr = run._r.get_or_add_rPr()
    ea = rPr.find(qn("a:ea"))
    if ea is None:
        ea = rPr.makeelement(qn("a:ea"), {}); rPr.append(ea)
    ea.set("typeface", name)


def rect(slide, x, y, w, h, fill=None, line=None, lw=1.0, shape=MSO_SHAPE.RECTANGLE):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid(); s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line; s.line.width = Pt(lw)
    s.shadow.inherit = False
    return s


def text(slide, x, y, w, h, s, size=16, bold=False, color=DARK, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, line=1.15, italic=False):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = Emu(0); tf.margin_right = Emu(0)
    tf.margin_top = Emu(0); tf.margin_bottom = Emu(0)
    p = tf.paragraphs[0]
    p.alignment = align
    p.line_spacing = line
    r = p.add_run(); r.text = s
    set_font(r, size, bold, color)
    if italic:
        r.font.italic = True
    return box


def header(slide, title, kicker=""):
    rect(slide, 0, 0, 13.333, 0.16, fill=NAVY)
    text(slide, 0.62, 0.34, 11.9, 0.62, title, 25, True, NAVY)
    if kicker:
        text(slide, 0.62, 0.98, 11.9, 0.3, kicker, 12, False, GRAY)
    rect(slide, 0.62, 1.30, 0.75, 0.045, fill=GOLD)


def footer(slide, idx):
    rect(slide, 0.62, 7.05, 12.1, 0.012, fill=RGBColor(0xD5, 0xDD, 0xE5))
    text(slide, 0.62, 7.12, 8.5, 0.28, "中国石油大学（北京） · 工商管理硕士（MBA）学位论文答辩", 9, False, GRAY)
    text(slide, 11.9, 7.12, 0.8, 0.28, str(idx), 10, False, GRAY, PP_ALIGN.RIGHT)


def bullets(slide, items, x, y, w, h, size=15, gap=0.14, color=DARK, lead_color=NAVY, line=1.2):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = Emu(0); tf.margin_right = Emu(0)
    tf.margin_top = Emu(0); tf.margin_bottom = Emu(0)
    first = True
    for it in items:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.line_spacing = line
        p.space_after = Pt(gap * 72)
        r0 = p.add_run(); r0.text = "▪ "
        set_font(r0, size, True, GOLD)
        if isinstance(it, tuple):
            lead, rest = it
            r1 = p.add_run(); r1.text = lead
            set_font(r1, size, True, lead_color)
            r2 = p.add_run(); r2.text = rest
            set_font(r2, size, False, color)
        else:
            r = p.add_run(); r.text = it
            set_font(r, size, False, color)
    return box


def pic(slide, key, x, y, maxw, maxh):
    path = os.path.join(FIG, FIGS[key])
    with Image.open(path) as im:
        iw, ih = im.size
    ratio = min(maxw / iw, maxh / ih)
    w, h = iw * ratio, ih * ratio
    cx = x + (maxw - w) / 2
    cy = y + (maxh - h) / 2
    slide.shapes.add_picture(path, Inches(cx), Inches(cy), Inches(w), Inches(h))
    return w, h


def table(slide, rows, x, y, w, col_ratios, size=11, row_h=0.42, header_fill=NAVY):
    ncol = len(rows[0]); nrow = len(rows)
    shape = slide.shapes.add_table(nrow, ncol, Inches(x), Inches(y), Inches(w), Inches(row_h * nrow))
    tbl = shape.table
    for i, r in enumerate(col_ratios):
        tbl.columns[i].width = Inches(w * r)
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            cell = tbl.cell(ri, ci)
            cell.text = ""
            tf = cell.text_frame
            tf.margin_left = Inches(0.08); tf.margin_right = Inches(0.06)
            tf.margin_top = Inches(0.03); tf.margin_bottom = Inches(0.03)
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.LEFT if ci == 0 else PP_ALIGN.CENTER
            run = p.add_run(); run.text = str(val)
            set_font(run, size, ri == 0, WHITE if ri == 0 else DARK)
            cell.fill.solid()
            cell.fill.fore_color.rgb = header_fill if ri == 0 else (LIGHT if ri % 2 == 0 else WHITE)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    return tbl


def notes(slide, s):
    slide.notes_slide.notes_text_frame.text = s


def cover():
    s = prs.slides.add_slide(BLANK)
    rect(s, 0, 0, 13.333, 7.5, fill=NAVY)
    rect(s, 0, 5.05, 13.333, 0.05, fill=GOLD)
    text(s, 0.9, 0.75, 8, 0.4, "中国石油大学（北京） · 经济管理学院", 14, False, RGBColor(0xBF, 0xD3, 0xE6))
    text(s, 0.9, 1.55, 11.6, 1.5, "Z公司工作流程数字化再造研究", 38, True, WHITE, line=1.15)
    text(s, 0.9, 3.05, 11.6, 0.6, "—— 基于招标代理业务线的环节级再造与平台实现", 18, False, RGBColor(0xCF, 0xDF, 0xEE))
    text(s, 0.9, 5.35, 5.6, 0.4, "答辩人：姚　烨　　学号：2024245834", 14, False, WHITE)
    text(s, 0.9, 5.85, 5.6, 0.4, "指导教师：郭庆方", 14, False, WHITE)
    text(s, 0.9, 6.35, 5.6, 0.4, "专业：工商管理（MBA）", 14, False, WHITE)
    text(s, 9.0, 6.35, 3.4, 0.4, "2026年10月", 14, False, RGBColor(0xCF, 0xDF, 0xEE), PP_ALIGN.RIGHT)
    notes(s, "各位老师上午好！我是2024级MBA姚烨，导师郭庆方老师。我的论文题目是《Z公司工作流程数字化再造研究》，聚焦招标代理业务线的环节级再造与平台实现。下面用约15分钟汇报。")
    return s

def toc_slide(title, left, right, idx):
    s = prs.slides.add_slide(BLANK)
    header(s, title, "共 7 章 + 引言、参考文献、致谢与附录；本页与下页为完整目录")
    def col(entries, x):
        y = 1.55
        for lvl, t in entries:
            if lvl == 0:
                text(s, x, y, 5.9, 0.3, t, 14, True, NAVY)
                rect(s, x, y + 0.33, 5.9, 0.008, fill=RGBColor(0xD5, 0xDD, 0xE5))
                y += 0.44
            elif lvl == 1:
                text(s, x + 0.18, y, 5.7, 0.26, t, 11.5, False, DARK)
                y += 0.265
            else:
                text(s, x + 0.5, y, 5.4, 0.24, t, 10.5, False, GRAY)
                y += 0.235
    col(left, 0.75); col(right, 6.85)
    footer(s, idx)
    notes(s, "这是论文的完整目录。第一部分是研究设计与问题诊断，第二部分是成因、方案、评估与结论。报告会按这个顺序展开。")
    return s

cover()

# ---------------- 目录（一）
toc_slide("目 录（一）｜研究设计与问题诊断", [
    (0, "引　言"),
    (0, "第1章　绪论"),
    (1, "1.1　研究背景与意义"),
    (1, "1.2　研究目的与方法"),
    (1, "1.3　文献综述"),
    (2, "1.3.1　业务流程再造与流程数字化研究"),
    (2, "1.3.2　招标采购数字化的实践与研究"),
    (2, "1.3.3　智能审核、风险识别与合规研究"),
    (2, "1.3.4　数字化转型的动力机制与理论框架"),
    (2, "1.3.5　研究述评与本文切入点"),
    (1, "1.4　本章小结"),
], [
    (0, "第2章　概念界定及相关理论"),
    (1, "2.1　概念界定"),
    (2, "2.1.1　工作流程数字化再造　　2.1.2　数字化转型"),
    (2, "2.1.3　业务流程集成与数字化成熟度"),
    (1, "2.2　相关理论"),
    (2, "2.2.1　业务流程再造理论　　2.2.2　TOE框架"),
    (2, "2.2.3　动态能力理论"),
    (1, "2.3　本章小结"),
    (0, "第3章　Z公司工作流程的现状与存在问题"),
    (1, "3.1　Z公司工作流程的现状"),
    (2, "3.1.1　数据来源与处理　　3.1.2　调查结果"),
    (1, "3.2　存在问题"),
    (2, "3.2.1　数字化工具集成度低"),
    (2, "3.2.2　流程节点冗余与时限管控薄弱"),
    (2, "3.2.3　员工数字化技能与工具采纳不足"),
    (1, "3.3　本章小结"),
], 2)

# ---------------- 目录（二）
toc_slide("目 录（二）｜成因、方案、评估与结论", [
    (0, "第4章　原因分析"),
    (1, "4.1　内部管理机制滞后"),
    (1, "4.2　技术基础设施薄弱"),
    (1, "4.3　流程设计与外部标准脱节"),
    (1, "4.4　人力资源培训体系不健全"),
    (1, "4.5　本章小结"),
    (0, "第5章　数字化工作活动平台的设计与实现"),
    (1, "5.1　平台总体设计"),
    (1, "5.2　流程建模与数据设计"),
    (1, "5.3　环节级流程再造设计"),
    (1, "5.4　关键功能模块实现"),
    (1, "5.5　系统集成、部署与实施推广"),
    (1, "5.6　本章小结"),
], [
    (0, "第6章　再造效果评估"),
    (1, "6.1　评价指标体系构建"),
    (1, "6.2　测量方法与数据可得性约束"),
    (1, "6.3　再造前后运行数据对比"),
    (1, "6.4　情景测算结果"),
    (1, "6.5　敏感性分析与讨论"),
    (1, "6.6　本章小结"),
    (0, "第7章　结论与展望"),
    (1, "7.1　研究结论　　7.2　研究局限与展望"),
    (0, "参考文献　致谢　附录A　附录B"),
    (1, "附录A　调查问卷（节选）　　附录B　访谈提纲"),
], 3)

def cards(slide, items, y=1.6, h=3.4, gap=0.28, x0=0.62, total_w=12.1):
    n = len(items)
    w = (total_w - gap * (n - 1)) / n
    for i, (title, lines, color) in enumerate(items):
        x = x0 + i * (w + gap)
        rect(slide, x, y, w, h, fill=RGBColor(0xF7, 0xFA, 0xFD), line=RGBColor(0xD5, 0xDD, 0xE5), lw=0.75)
        rect(slide, x, y, w, 0.055, fill=color)
        text(slide, x + 0.22, y + 0.2, w - 0.44, 0.36, title, 15, True, NAVY)
        yy = y + 0.68
        for ln in lines:
            text(slide, x + 0.22, yy, w - 0.44, 0.52, ln, 11.5, False, DARK, line=1.25)
            yy += 0.58 if len(ln) < 20 else 0.78

# ---------------- S4 研究背景
s = prs.slides.add_slide(BLANK)
header(s, "研究背景：三股力量把“流程”推到前台", "为什么研究招标代理的流程数字化再造？")
cards(s, [
 ("政策之势", ["数字经济成为国家战略", "数据要素市场化配置加速", "企业从经验驱动转向数据驱动"], BLUE),
 ("行业之势", ["2013年电子招投标办法", "全流程电子化持续深化", "2024年公平竞争审查规则强化事前合规"], NAVY),
 ("竞争之势", ["服务收费实行市场调节价", "机构同质化竞争激烈", "向“效率+合规+专业”要效益"], GOLD),
], y=1.62, h=3.25)
rect(s, 0.62, 5.1, 12.1, 1.55, fill=RGBColor(0x12, 0x3A, 0x63))
text(s, 0.9, 5.28, 11.6, 0.4, "企业之困：平台上线 ≠ 流程再造", 17, True, WHITE)
text(s, 0.9, 5.75, 11.6, 0.8,
     "Z公司2025年2月上线智能作业平台（V1），但部门系统各自为政、节点时限缺少统一度量、成果质量依赖个人经验——"
     "“有系统”不等于“流程被重构”，这正是本文要回答的问题。", 12.5, False, RGBColor(0xE6, 0xEE, 0xF6), line=1.3)
footer(s, 4)
notes(s, "先看背景：政策、行业、竞争三股力量叠加，招标代理的流程必须从“人工经验”转向“平台化、可度量”。但Z公司2025年2月上线平台后，仍然存在系统各自为政、节点时限不可度量的问题，说明上线不等于再造。")

# ---------------- S5 研究问题
s = prs.slides.add_slide(BLANK)
header(s, "研究问题：一个“诊断—设计—评估”闭环", "四个问题，对应论文的四段主线")
qs = [
 ("Q1　现状与问题", "招标代理流程的数字化到了什么程度？卡在哪三个环节？"),
 ("Q2　成因", "为什么会卡？管理、技术、环境、能力各自起什么作用？"),
 ("Q3　方案", "如何把“再造”落到可运行的平台与可测量的数据？"),
 ("Q4　效果", "再造是否改善了效率与合规？在数据受限下如何评估？"),
]
for i, (q, a) in enumerate(qs):
    x = 0.62 + (i % 2) * 6.2
    y = 1.65 + (i // 2) * 2.15
    rect(s, x, y, 5.9, 1.85, fill=RGBColor(0xF7, 0xFA, 0xFD), line=RGBColor(0xD5, 0xDD, 0xE5), lw=0.75)
    rect(s, x, y, 0.07, 1.85, fill=GOLD)
    text(s, x + 0.28, y + 0.2, 5.4, 0.35, q, 16, True, NAVY)
    text(s, x + 0.28, y + 0.72, 5.35, 0.95, a, 12.5, False, DARK, line=1.3)
text(s, 0.62, 6.0, 12.1, 0.5, "研究目标：形成“问题诊断 → 环节再造 → 平台实现 → 效果评估”的完整闭环，为同类机构提供可复用的方法与工具。",
     13.5, True, NAVY)
footer(s, 5)
notes(s, "围绕四个问题：现状卡在哪、为什么卡、怎么改、改得怎么样。对应论文的“诊断—设计—评估”闭环。")

# ---------------- S6 研究方法与技术路线
s = prs.slides.add_slide(BLANK)
header(s, "研究方法：多源证据的混合方法", "单案例 + 问卷 + 访谈 + 设计科学 + 情景分析")
bullets(s, [
 ("文献研究：", "52篇中英文文献，梳理BPR、TOE、动态能力与招投标数字化。"),
 ("案例研究：", "Z公司招标代理业务线，嵌入多源证据（问卷、访谈、台账、系统文档）。"),
 ("问卷调查：", "2024年10—11月，发放200份、有效185份（92.5%），25题（7基本+13 Likert+5开放）。"),
 ("半结构化访谈：", "10名中高层与资深项目经理，平均45分钟，主题编码。"),
 ("设计科学研究：", "需求→设计→实现→验证，产出可运行平台。"),
 ("情景分析：", "保守/基准/乐观三情景，覆盖无实测基线的指标。"),
], 0.62, 1.55, 6.2, 4.9, size=13, gap=0.12)
pic(s, "route", 6.95, 1.5, 5.8, 5.0)
footer(s, 6)
notes(s, "方法上是单案例嵌入混合方法：问卷185份、访谈10人，再叠加业务台账与系统文档做三角验证；方案部分用设计科学研究的路径，效果部分用情景分析补足无基线指标。")

# ---------------- S7 理论基础
s = prs.slides.add_slide(BLANK)
header(s, "理论基础：三个理论各管一段", "BPR 管“怎么改”，TOE 管“为什么能/不能改”，动态能力管“改完能否持续”")
cards(s, [
 ("业务流程再造（BPR）", ["Hammer：根本性反思、彻底性再设计", "ESIA：清除、简化、整合、自动化", "本文用途：环节级再造的理论依据"], NAVY),
 ("TOE 框架", ["技术—组织—环境三维度", "解释数字化采纳的条件与障碍", "本文用途：四类成因的归位框架"], BLUE),
 ("动态能力理论", ["感知—抓取—重构", "解释企业持续适应环境的能力", "本文用途：解释“可测量性”与持续优化"], GOLD),
], y=1.65, h=3.0)
text(s, 0.62, 4.95, 12.1, 0.45, "整合逻辑：BPR 提供再造方法，TOE 解释约束条件，动态能力解释长期适应；三者共同支撑“诊断—设计—评估”框架。",
     12.5, False, DARK)
pic(s, "framework", 0.62, 4.55, 12.1, 2.3)
footer(s, 7)
notes(s, "理论上不是贴标签：BPR 告诉我改什么（环节的清除/简化/整合/自动化），TOE 告诉我为什么改不动（技术、组织、环境），动态能力解释改完之后怎么持续。三者整合成图2.1的框架。")

# ---------------- S8 文献综述
s = prs.slides.add_slide(BLANK)
header(s, "文献综述：五条线索，一个缺口", "已有研究多停在“技术可行”或“制度环境”，缺少“环节级再造 + 可测量评估”的整合")
bullets(s, [
 ("流程再造与流程数字化：", "从 Hammer 的彻底重构，到国内学者对流程数字化的本土化研究。"),
 ("招标采购数字化：", "电子招投标、清标自动化、围串标识别的实践与研究。"),
 ("智能审核与合规：", "NLP、规则引擎、知识库在合规审查中的应用。"),
 ("转型动力机制：", "TOE、资源基础观、动态能力对数字化采纳的解释。"),
 ("研究缺口：", "① 多聚焦单点技术，少谈环节重组；② 效果评估多为定性，缺可测量口径；③ 数据受限场景下的评估方法缺失。"),
], 0.62, 1.6, 12.1, 4.2, size=13.5, gap=0.16)
rect(s, 0.62, 5.55, 12.1, 1.05, fill=RGBColor(0xEA, 0xF1, 0xF8))
rect(s, 0.62, 5.55, 0.08, 1.05, fill=GOLD)
text(s, 0.95, 5.7, 11.6, 0.75, "本文切入点：把 BPR 的“环节级再造”落到系统设计，把效果评估落到“可测量指标 + 数据受限下的组合证据”。",
     13.5, True, NAVY, line=1.25)
footer(s, 8)
notes(s, "文献部分归纳了五条线索，最后落到三个缺口：单点技术多、环节重组少；定性多、可测量少；数据受限下的评估方法缺失。本文的切入点就是把环节再造和可测量评估接起来。")

# ---------------- S9 概念界定
s = prs.slides.add_slide(BLANK)
header(s, "核心概念：三个关键词的边界", "避免把“数字化”“信息化”“流程再造”混为一谈")
cards(s, [
 ("工作流程数字化再造", ["以数字技术为杠杆，对流程环节做根本性再设计", "不是把线下流程原样搬到线上", "判据：环节是否被清除/合并/并行/自动化"], NAVY),
 ("数字化转型", ["技术—业务—组织协同的系统变革", "强调数据要素与业务场景融合", "不止是工具替换，更是能力重建"], BLUE),
 ("业务流程集成与成熟度", ["工具互联、数据贯通、流程可计算", "成熟度：从“事后统计”到“实时可见”", "本文用它衡量“集成水平”维度"], GOLD),
], y=1.65, h=3.4)
text(s, 0.62, 5.45, 12.1, 0.6, "本文的界定：工作流程数字化再造 = 环节级流程重构 + 平台化承载 + 可测量治理。",
     14, True, NAVY)
footer(s, 9)
notes(s, "概念上先划边界：再造不是把线下流程搬到线上，判据是环节有没有被清除、合并、并行或自动化；数字化转型是技术、业务、组织协同的变革；集成与成熟度则用来衡量“可计算”的程度。")

# ---------------- S10 数据来源与样本
s = prs.slides.add_slide(BLANK)
header(s, "数据与样本：四类证据相互印证", "2024年10—11月调研，处于平台上线（2025-02）之前，反映再造前状况")
bullets(s, [
 ("问卷：", "200份发放、185份有效（92.5%）；25题，覆盖基本信息、流程现状、数字化应用、能力与问题感知。"),
 ("访谈：", "10名中高层与资深项目经理，平均45分钟，按“工具与集成/流程与时限/人员与培训”编码。"),
 ("业务台账：", "再造前12个月约450个项目、平均周期约27天；再造后约570个、约20天（均含等待与澄清期）。"),
 ("系统资料：", "制度文件、平台设计文档、数据库表结构、接口与流程日志。"),
], 0.62, 1.6, 6.2, 4.2, size=13, gap=0.14)
pic(s, "sample", 7.0, 1.5, 5.7, 4.9)
footer(s, 10)
notes(s, "数据分四类：问卷、访谈、业务台账、系统资料。特别强调：问卷是2024年10—11月做的，在平台上线之前，所以它反映的是再造前的状况，这和第6章的“再造前12个月”是同一条基线。")

# ---------------- S11 六阶段流程
s = prs.slides.add_slide(BLANK)
header(s, "现状扫描：六阶段流程与三个瓶颈", "流程框架是完整的，问题出在“衔接”与“时限”")
steps = ["受理立项", "文件编制", "公告答疑", "开标清标", "评审定标", "归档合同"]
x = 0.62
for i, st in enumerate(steps):
    rect(s, x, 1.65, 1.85, 0.62, fill=NAVY if i % 2 == 0 else BLUE)
    text(s, x, 1.65, 1.85, 0.62, st, 12, True, WHITE, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
    if i < 5:
        text(s, x + 1.87, 1.7, 0.25, 0.4, "›", 20, True, GOLD)
    x += 2.02
bullets(s, [
 ("瓶颈一：衔接靠人工。", "从文件编制到合规核对、从清标到评审报告，依赖人工传递，访谈估计平均耗时3—5天。"),
 ("瓶颈二：复核层级多。", "成果复核多轮审核，每轮重新审阅与修改，平均约2天延时。"),
 ("瓶颈三：纸质与线下。", "40.0%的受访者认为纸质流转频繁，打印、签字、存档造成额外消耗。"),
 ("本质：", "流程框架标准化程度不低，但“环节之间的接口”没有被系统承载，时限也无法度量。"),
], 0.62, 2.55, 12.1, 3.9, size=13, gap=0.14)
footer(s, 11)
notes(s, "六阶段流程本身是完整的，标准化程度也不低（85.4%的受访者认可），但三个瓶颈都出在“衔接”：人工传递3—5天、多轮复核约2天、纸质流转40%。这说明问题不在流程框架，而在环节之间的接口没有被系统承载。")

# ---------------- S12 问题诊断
s = prs.slides.add_slide(BLANK)
header(s, "问题诊断：三个问题，指向同一个“不可计算”", "问卷（N=185）与访谈相互印证")
bullets(s, [
 ("① 集成度低：", "仅45.9%认同工具间无缝对接；数据孤岛、重复录入、口径不一。"),
 ("② 节点冗余、时限弱：", "52.4%认为节点衔接不畅；多轮复核约2天；纸质流转40.0%。"),
 ("③ 技能与采纳不足：", "55.0%认为培训覆盖不足；数字化意识得分3.7/5，工具采纳有限。"),
 ("访谈原话：", "“投标文件解析在一个系统，清标结果又要在另一个表格里汇总，来回导数据很容易出错。”"),
], 0.62, 1.6, 6.1, 4.6, size=13, gap=0.16)
pic(s, "problem", 6.9, 1.5, 5.8, 5.0)
footer(s, 12)
notes(s, "三个问题：集成度低（45.9%）、节点冗余时限弱（52.4%、复核2天、纸质40%）、技能不足（55%）。访谈原话很典型：解析在一个系统、清标在另一个表格，来回导数据。三个问题的共同点是——流程不可计算。")

# ---------------- S13 证据表
s = prs.slides.add_slide(BLANK)
header(s, "问题证据：把“感受”变成“可核对的事实”", "每个问题都对应可追溯的数据来源")
rows = [
 ["问题", "关键证据（2024年调研，N=185）", "管理含义"],
 ["数字化工具集成度低", "45.9%认同工具间无缝对接；数据孤岛、重复录入", "数据口径不一，跨部门协同成本高"],
 ["流程节点冗余、时限弱", "52.4%认为衔接不畅；复核约2天；人工传递3—5天", "周期被非增值等待拉长，返工增加"],
 ["员工技能与工具采纳不足", "55.0%认为培训覆盖不足；数字化意识3.7/5", "工具闲置，质量依赖个人经验"],
]
table(s, rows, 0.62, 1.7, 12.1, [0.22, 0.46, 0.32], size=12, row_h=0.62)
text(s, 0.62, 4.75, 12.1, 0.9, "说明：百分比为问卷分项统计结果；3—5天为访谈估计；复核延时取“约2天”。上述证据与第6章“再造前12个月”台账基线相互印证。",
     11, False, GRAY, line=1.3)
footer(s, 13)
notes(s, "这张表把三个问题、证据、管理含义对齐，方便评委核对。数据都来自2024年的问卷、访谈和台账，和后面第6章的前后对比是同一套口径。")

# ---------------- S14 成因
s = prs.slides.add_slide(BLANK)
header(s, "成因分析：四条成因，一条链", "不是四个并列原因，而是相互强化的结构性约束")
bullets(s, [
 ("管理机制滞后：", "决策链条长、部门协调不畅、绩效偏重业务量。"),
 ("技术基础薄弱：", "系统分散、接口标准缺失、算力与存储不足。"),
 ("流程与外部标准脱节：", "法规与数据接口快速迭代，流程未同步精简。"),
 ("培训体系不健全：", "课程结构失衡、覆盖不均，缺少流程再造思维训练。"),
], 0.62, 1.6, 6.1, 4.6, size=13, gap=0.16)
pic(s, "fishbone", 6.85, 1.5, 5.9, 5.0)
footer(s, 14)
notes(s, "四条成因：管理机制、技术基础、外部标准脱节、培训体系。用鱼骨图是为了说明它们都指向三个问题，而不是各管一段。")

# ---------------- S15 传导机制
s = prs.slides.add_slide(BLANK)
header(s, "传导机制：为什么“单点优化”无效", "问题—成因—对策的链条决定了必须“环节 + 平台 + 组织”三位一体")
pic(s, "chain", 0.62, 1.5, 8.15, 4.6)
bullets(s, [
 ("相互强化：", "管理制约技术投入，技术放大流程冗余，外部变化又放大能力短板。"),
 ("结论一：", "问题是系统性的，单点上线工具无法解决。"),
 ("结论二：", "对策必须同时动“环节”（ESIA）、“平台”（可计算）、“组织”（绩效与培训）。"),
], 8.95, 1.6, 3.8, 4.6, size=12, gap=0.16)
footer(s, 15)
notes(s, "这条传导链说明：四类成因相互强化，形成结构性约束。所以单点优化无效，必须环节、平台、组织三位一体——这也直接引出第5章的设计。")

# ---------------- S16 平台总体设计
s = prs.slides.add_slide(BLANK)
header(s, "方案总体设计：五项建设目标", "从“看得见”到“算得出”，再到“管得住”")
bullets(s, [
 ("流程可见：", "项目—时间轴—里程碑—环节—待办对象化，进度、责任人、时限实时可见。"),
 ("作业可助：", "文档解析、清标比对、条款审查等重复劳动由系统辅助。"),
 ("风险可控：", "规则库与风险模型，把合规与围串标识别前置。"),
 ("知识可积：", "法规、范本、历史成果沉淀为可检索的组织知识。"),
 ("环节精简：", "按ESIA原则删除或合并非增值环节，环节净减、责任清晰。"),
], 0.62, 1.55, 5.9, 4.4, size=12.5, gap=0.13)
pic(s, "arch", 6.6, 1.45, 6.2, 5.1)
footer(s, 16)
notes(s, "平台设计五个目标：流程可见、作业可助、风险可控、知识可积、环节精简。前四个是“让系统帮忙”，第五个是“动流程本身”——这是本文和单纯做系统最大的区别。")

# ---------------- S17 五级对象模型
s = prs.slides.add_slide(BLANK)
header(s, "五级流程对象：把流程变成可计算的数据", "项目—时间轴—里程碑—环节—待办")
pic(s, "data", 0.62, 1.5, 7.5, 4.9)
bullets(s, [
 ("对象化：", "五级对象承载全周期，责任、时限、状态都可查询。"),
 ("可计算：", "流程周期、计划偏差率、里程碑延误、返工强度、任务响应。"),
 ("可追溯：", "文件版本化、结果快照化、法规版本化。"),
 ("数据治理：", "一次录入、多处复用，减少重复录入与口径不一。"),
], 8.35, 1.6, 4.35, 4.6, size=12, gap=0.16)
footer(s, 17)
notes(s, "这是平台的数据模型：五级对象把流程拆成可计算的结构，于是“周期、偏差、延误、返工、响应”这些指标才能自动算出来。没有这层对象化，后面的效果评估就无从谈起。")

# ---------------- S18 ESIA
s = prs.slides.add_slide(BLANK)
header(s, "环节级再造：ESIA 五刀，刀刀落在冗余点", "这是“再造”区别于“上线”的关键：动的是环节，不是界面")
rows = [
 ["冗余环节（第3章证据）", "ESIA 再造设计", "实施状态"],
 ["复核三级会签（平均延时约2天）", "清除重复会签：三级复核改为“一级复核 + 异常抽检”", "已实施"],
 ["纸质签章与归档（40.0%认为纸质频繁）", "自动化与简化：电子签章、电子档案、归档线上化", "已实施"],
 ["串行会签、多方等待", "整合与并行：线上并行会签、超时提醒", "已实施"],
 ["清标逐项人工比对", "自动化：系统五维初筛 + 人工复核", "已实施"],
 ["人工台账跟踪节点", "简化：时间轴与待办自动记录节点状态", "已实施"],
]
table(s, rows, 0.62, 1.65, 12.1, [0.34, 0.50, 0.16], size=11.5, row_h=0.62)
text(s, 0.62, 5.55, 12.1, 0.8, "效果：复核层级由三级降为一级并对异常抽检；清标由“全量人工比对”转为“异常复核”；签章与会签线上化，纸质流转显著减少。",
     12, False, DARK, line=1.3)
footer(s, 18)
notes(s, "这一页是全文的关键：ESIA五刀——复核降级、签章电子化、会签并行、清标系统初筛、台账自动化，而且都已经在业务线实施。这是“再造”与“上线”的分水岭。")

# ---------------- S19 功能模块
s = prs.slides.add_slide(BLANK)
header(s, "关键功能模块：七个模块支撑一条流程", "模块不是堆功能，每个都对应第3章的一个问题")
rows = [
 ["模块", "核心能力", "对应的管理问题"],
 ["项目管理与时间轴", "五级对象、进度视图、节点预警", "流程不可见、时限不可度量"],
 ["文档智能解析", "多格式解析、条款提取、结构化", "人工通读耗时、口径不一"],
 ["清标辅助评分", "五维复合评分、交叉比较、异常提示", "逐项人工比对、易出错"],
 ["围串标风险识别", "相似度、尾数、Benford 检测", "依赖个人经验、识别滞后"],
 ["合规检查", "15部法规、619条规则库", "条款审查慢、标准不统一"],
 ["智能对话与知识库", "大模型问答、RAG、法规版本", "知识沉淀不足、重复问答"],
 ["项目协同与信用查询", "并行会签、企业信用信息", "协同靠线下、信息不全"],
]
table(s, rows, 0.62, 1.6, 12.1, [0.24, 0.42, 0.34], size=10.5, row_h=0.52)
text(s, 0.62, 5.9, 12.1, 0.5, "技术底座：Flask 3.1 + PostgreSQL 16 + Redis 7 + Celery + LangGraph + ChromaDB，本地化部署大模型，数据不出内网。",
     10.5, False, GRAY)
footer(s, 19)
notes(s, "七个功能模块，每个都对应前面诊断出的问题：时间轴对“不可见”，解析和清标对“人工耗时”，围串标和合规对“依赖经验”，知识库对“知识不沉淀”。技术上是私有化部署，数据不出内网。")

# ---------------- S20 实施推广
s = prs.slides.add_slide(BLANK)
header(s, "实施推广：四步走 + 四项组织配套", "系统上线只是开始，配套不同步会稀释收益")
path = [("试点", "两个业务组验证"), ("评估", "节点耗时与返工"), ("推广", "业务线全面上线"), ("固化", "制度与绩效")]
x = 0.62
for i, (t, d) in enumerate(path):
    rect(s, x, 1.62, 2.7, 0.75, fill=NAVY if i % 2 == 0 else BLUE)
    text(s, x, 1.62, 2.7, 0.75, t, 15, True, WHITE, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
    text(s, x, 2.45, 2.7, 0.3, d, 10.5, False, GRAY, PP_ALIGN.CENTER)
    if i < 3:
        text(s, x + 2.72, 1.75, 0.35, 0.4, "→", 18, True, GOLD)
    x += 3.05
bullets(s, [
 ("流程治理：", "成立流程治理小组，明确节点责任人与时限。"),
 ("绩效配套：", "把节点准时率、一次通过率纳入考核，弱化单纯业务量导向。"),
 ("能力配套：", "分层培训（管理者/项目经理/一般员工），培训覆盖率纳入验收。"),
 ("制度固化：", "把流程节点、报送要求、留痕标准写入制度，避免“一阵风”。"),
], 0.62, 3.15, 12.1, 3.0, size=13, gap=0.16)
footer(s, 20)
notes(s, "实施路径是试点—评估—推广—固化，同时配四件事：流程治理、绩效、培训、制度。论文里反复强调：系统上线不等于再造，配套不同步，收益会被稀释。")

# ---------------- S21 评价体系
s = prs.slides.add_slide(BLANK)
header(s, "效果评估：5个准则、17项指标", "先建口径，再谈结果——数据受限下的评估设计")
bullets(s, [
 ("五个准则：", "流程效率、集成水平、资源节约、业务价值、组织能力。"),
 ("数据约束：", "只有聚合数据、没有逐项目明细，因此不做显著性检验。"),
 ("证据分层：", "3项前后实测（项目数、周期、清标）+ 4项上线后观测（接口成功率、自动流转、月活、培训覆盖）+ 10项待验证。"),
 ("评估逻辑：", "能实测的用实测，无基线的用情景，无法量化的只作定性说明。"),
], 0.62, 1.6, 6.1, 4.6, size=13, gap=0.16)
pic(s, "kpi", 6.9, 1.5, 5.8, 5.0)
footer(s, 21)
notes(s, "评估体系5个准则17项指标，但受数据限制，真正有前后实测的只有3项，4项是上线后单点观测，其余10项只建了口径。这个“证据分层”是论文的诚实之处，也是答辩时最容易被问的地方。")

# ---------------- S22 实测结果
s = prs.slides.add_slide(BLANK)
header(s, "实测结果：周期与清标显著改善", "再造前后各12个月，经公司许可的聚合数据")
pic(s, "measured", 0.62, 1.5, 7.4, 3.7)
bullets(s, [
 ("项目数：", "约450 → 约570个（约 +27%）"),
 ("流程周期：", "约27 → 约20天（约 −26%）"),
 ("清标耗时：", "约3个工作日 → 约20分钟"),
 ("其他观测：", "接口成功率 >99.9%、自动流转约75%、月活与培训覆盖100%（单点，无基线）"),
], 8.25, 1.55, 4.5, 3.6, size=12, gap=0.14)
rect(s, 0.62, 5.55, 12.1, 1.05, fill=RGBColor(0xEA, 0xF1, 0xF8))
rect(s, 0.62, 5.55, 0.08, 1.05, fill=GOLD)
text(s, 0.95, 5.68, 11.6, 0.8, "口径与归因：上线前取业务台账（含等待与澄清期），上线后取平台时间轴；变化与平台上线同期发生，受业务量与结构变化影响，尚不能单独归因。",
     11.5, True, NAVY, line=1.3)
footer(s, 22)
notes(s, "实测结果三个数：项目数450到570、周期27到20天、清标3个工作日到20分钟。注意口径：上线前台账、上线后时间轴，都含等待期；而且这是同期变化，不能单独归因于平台，这一点论文里写得很清楚。")

# ---------------- S23 返工与情景
s = prs.slides.add_slide(BLANK)
header(s, "返工与情景：有观测的看观测，没基线的看区间", "测算值不等于实测值，结论强度要匹配证据强度")
pic(s, "rework", 0.62, 1.5, 5.7, 3.4)
pic(s, "scenario", 6.75, 1.5, 5.95, 3.4)
text(s, 0.62, 5.05, 5.7, 0.3, "图：分业务类型平均返工轮次（上线后实测）", 10, False, GRAY)
text(s, 6.75, 5.05, 5.95, 0.3, "图：三情景指标变化与综合指数（测算）", 10, False, GRAY)
bullets(s, [
 ("返工：", "总体1.8次；一般1.2 / 政采2.2 / 部队及武警3.0；人工搬运5次/项目。"),
 ("情景：", "保守/基准/乐观三情景；综合指数100→67.0（82.3 / 67.0 / 54.5）。"),
 ("边界：", "参数未校准（周期实测−26%已超乐观−22%），情景值仅作方向性参考，综合指数仅作方法演示。"),
], 0.62, 5.45, 12.1, 1.4, size=11.5, gap=0.1)
footer(s, 23)
notes(s, "返工和搬运没有再造前基线，只能给情景区间；综合指数67.0只是方法演示。特别说明：实测周期降幅已经超过乐观情景，说明参数没校准，所以情景值只能看方向，不能当结论。")

# ---------------- S24 局限与结论
s = prs.slides.add_slide(BLANK)
header(s, "研究结论、贡献与局限", "结论不放大，局限不遮掩")
rect(s, 0.62, 1.6, 5.85, 4.55, fill=RGBColor(0xF7, 0xFA, 0xFD), line=RGBColor(0xD5, 0xDD, 0xE5), lw=0.75)
rect(s, 0.62, 1.6, 5.85, 0.06, fill=GOLD)
text(s, 0.85, 1.78, 5.4, 0.35, "研究结论", 15, True, NAVY)
bullets(s, [
 ("问题：", "三个问题共同指向“流程不可计算”，且由四类成因交互强化。"),
 ("方案：", "ESIA环节再造 + 五级对象平台 + 组织配套，三位一体。"),
 ("效果：", "流程周期与清标环节同期显著改善；集成与组织维度达到较高运行水平。"),
 ("方法：", "形成“可测量指标 + 数据受限下组合证据”的评估路径。"),
], 0.85, 2.2, 5.35, 3.6, size=11.5, gap=0.1, lead_color=BLUE)
rect(s, 6.7, 1.6, 6.0, 4.55, fill=RGBColor(0xFD, 0xF8, 0xEC), line=RGBColor(0xE8, 0xD9, 0xB0), lw=0.75)
rect(s, 6.7, 1.6, 6.0, 0.06, fill=GOLD)
text(s, 6.93, 1.78, 5.5, 0.35, "主要局限", 15, True, NAVY)
bullets(s, [
 ("数据：", "仅聚合数据，无逐项目明细，未做统计检验。"),
 ("基线：", "返工、搬运缺再造前口径，只能给情景区间。"),
 ("范围：", "2026年V2大模型能力未纳入本次实测。"),
 ("混杂：", "业务量增长约27%、结构变化可能影响周期对比。"),
 ("口径：", "清标前后计时口径不完全一致，仅作方向性比较。"),
], 6.93, 2.2, 5.5, 3.7, size=11.5, gap=0.08, lead_color=BLUE)
footer(s, 24)
notes(s, "结论三句话：问题系统性、方案三位一体、周期与清标同期改善。局限也直说：聚合数据、无检验、V2未纳入、业务量混杂、清标口径不一致。结论强度与证据强度匹配，不放大。")

# ---------------- S25 致谢
s = prs.slides.add_slide(BLANK)
rect(s, 0, 0, 13.333, 7.5, fill=NAVY)
rect(s, 0, 5.2, 13.333, 0.05, fill=GOLD)
text(s, 0.9, 1.35, 11.5, 0.9, "汇报完毕，恳请各位老师批评指正", 30, True, WHITE)
text(s, 0.9, 2.45, 11.5, 0.5, "感谢导师郭庆方教授的悉心指导，感谢Z公司提供的调研与数据支持，感谢各位评委老师。", 14, False, RGBColor(0xCF, 0xDF, 0xEE), line=1.4)
text(s, 0.9, 5.55, 6.0, 0.4, "答辩人：姚　烨", 13, False, WHITE)
text(s, 0.9, 6.0, 6.0, 0.4, "学号：2024245834　｜　工商管理（MBA）", 13, False, WHITE)
text(s, 7.6, 6.0, 4.8, 0.4, "中国石油大学（北京）· 经济管理学院", 12, False, RGBColor(0xBF, 0xD3, 0xE6), PP_ALIGN.RIGHT)
footer(s, 25)
notes(s, "我的汇报到此结束，恳请各位老师批评指正，谢谢！")

out = "v12/Z公司工作流程数字化再造研究_答辩PPT.pptx"
prs.save(out)
print("saved", out, "| slides:", len(prs.slides.__iter__.__self__._sldIdLst))
