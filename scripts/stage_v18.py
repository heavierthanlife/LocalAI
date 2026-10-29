# -*- coding: utf-8 -*-
"""v1.8 阶段一：按"问卷在 2024-10/11（上线前）"重顺因果链；补回附录A/B分节；恢复自动目录域。"""
import copy
import docx
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph

SRC = "v12/user_v11414.docx"
OUT = "v12/_v18_stage.docx"

d = docx.Document(SRC)
ps = d.paragraphs

# ---------------- 文本修订（因果链） ----------------
REPL = [
 # 1) 3.1 现状定界段
 ("需要说明的是，受数据可得性限制，问卷与访谈于2025年10—11月开展，反映的是平台上线后仍存在的残余问题，因此第3章的问题诊断针对“再造后仍需改进的环节”，而非再造前的初始状态。",
  "需要说明的是，问卷与访谈于2024年10—11月开展，处于平台上线之前，反映的是再造前的流程状况，与业务台账的“前”侧基线相互印证；因此第3章的问题诊断针对再造前存在的突出问题，构成后续再造方案设计的事实基础。"),
 # 2) 3.1.1 问卷时间
 ("调查于2025年10月至11月进行，通过公司内部OA系统和微信平台发放",
  "调查于2024年10月至11月进行，通过公司内部OA系统和微信平台发放"),
]
for old, new in REPL:
    hit = 0
    for p in ps:
        if old in p.text:
            # 段落内替换：尽量保留 run 结构
            replaced = False
            for r in p.runs:
                if old in r.text:
                    r.text = r.text.replace(old, new)
                    replaced = True
                    break
            if not replaced:
                p.text = p.text.replace(old, new)
            hit += 1
    print(f"[文本替换] {'OK' if hit else 'MISS'} ×{hit}: {old[:24]}…")

# 3) 3.1 数字化段：调研时点（2024）与上线（2025-02）分开
OLD3 = ("在数字化方面，招标代理业务线已于2025年2月上线智能作业平台（V1），2026年重构为V2并引入大模型能力；"
        "两版平台共用同一套流程对象与数据模型，V2相对V1增加了本地化大模型驱动的内容解析、合规检查与知识检索能力。")
NEW3 = ("在数字化方面，调研时（2024年）公司已引入文档解析、清标辅助等数字化工具，但相关能力分散在多个系统，尚未形成统一的流程平台；"
        "2025年2月，公司在此基础上上线智能作业平台（V1），作为本文的再造分界；2026年重构为V2并引入本地化大模型能力，"
        "两版平台共用同一套流程对象与数据模型。")
n = 0
for p in ps:
    if OLD3 in p.text:
        for r in p.runs:
            if OLD3 in r.text:
                r.text = r.text.replace(OLD3, NEW3); n += 1; break
        else:
            p.text = p.text.replace(OLD3, NEW3); n += 1
print("[3.1 数字化段]", "OK" if n else "MISS", n)

# 4) 3.1.2 数字化应用段：调研时点与整合上线分开
OLD4 = "在数字化应用方面，公司已引入多项工具。2025年2月上线的智能作业平台支持文档全格式解析、清标辅助评分、合规检查与知识库检索，实现部分环节自动化。"
NEW4 = ("在数字化应用方面，调研时公司已引入多项数字化工具，文档全格式解析、清标辅助评分、合规检查与知识库检索等能力分散在不同系统，部分环节实现自动化；"
        "2025年2月，上述能力整合为统一的智能作业平台（V1）。")
n = 0
for p in ps:
    if OLD4 in p.text:
        for r in p.runs:
            if OLD4 in r.text:
                r.text = r.text.replace(OLD4, NEW4); n += 1; break
        else:
            p.text = p.text.replace(OLD4, NEW4); n += 1
print("[3.1.2 数字化应用段]", "OK" if n else "MISS", n)

# ---------------- 结构：补回 附录A / 附录B 分节 ----------------
target = None
for p in ps:
    if p.style.name == "标题1" and p.text.strip().startswith("附录B"):
        target = p
        break
assert target is not None, "未找到附录B标题"
sectPr = copy.deepcopy(d.sections[-1]._sectPr)
newp = OxmlElement("w:p")
pPr = OxmlElement("w:pPr")
pPr.append(sectPr)
newp.append(pPr)
target._p.addprevious(newp)
print("[分节] 已在附录B前补分节；当前节数:", len(d.sections))

# ---------------- 目录：静态条目 → 自动域 ----------------
toc_ps = [p for p in d.paragraphs if p.style.name.lower().startswith("toc")]
if toc_ps:
    first = toc_ps[0]._p
    holder = OxmlElement("w:p")
    r = OxmlElement("w:r")
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), 'TOC \\o "1-3" \\h \\z \\u')
    ph = OxmlElement("w:r"); rPr = OxmlElement("w:rPr")
    col = OxmlElement("w:color"); col.set(qn("w:val"), "808080")
    it = OxmlElement("w:i"); rPr.append(it); rPr.append(col); ph.append(rPr)
    t = OxmlElement("w:t"); t.text = "（目录为自动域：右键 → 更新域 → 更新整个目录）"
    ph.append(t); fld.append(ph); r.append(fld); holder.append(r)
    first.addprevious(holder)
    for p in toc_ps:
        p._p.getparent().remove(p._p)
    print("[目录] 已替换静态目录为自动域，移除条目:", len(toc_ps))
else:
    print("[目录] 未发现静态目录")

d.save(OUT)
print("saved", OUT, "| 段落", len(docx.Document(OUT).paragraphs), "| 节", len(docx.Document(OUT).sections))
