# -*- coding: utf-8 -*-
"""生成 700 字新摘要（单独 docx，按《格式.doc》摘要格式）。"""
import re
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

SONG, HEI, TNR, ARIAL = "宋体", "黑体", "Times New Roman", "Arial"

PARAS = [
"在数字经济与招标投标全流程电子化持续深化的背景下，招标代理机构的工作流程数字化再造已成为提升效率与合规水平的重要路径。招标代理业务文档密集、合规要求高，传统人工审核、纸质流转与串行会签难以适应监管迭代与客户时效要求。本研究以Z公司招标代理业务线为对象，探讨其流程数字化再造的理论框架、实现路径与效果评价。",
"研究综合运用文献综述、案例研究、问卷调查、半结构化访谈、设计科学研究与情景分析等方法，将问卷、访谈、制度文件、系统设计文档与经公司许可的聚合数据相互印证，形成混合方法分析框架。",
"研究发现，Z公司流程已初步数字化，但仍存在三方面问题：工具间缺乏有效衔接，数据孤岛与重复录入并存；流程节点冗余、时限管控薄弱，复核与纸质流转消耗较多非增值时间；员工数字化意识与操作技能不足，工具采纳有限。成因在于管理机制滞后、技术基础薄弱、流程设计与外部标准脱节、培训体系不健全，四者交互并相互强化。",
"基于业务流程再造理论、TOE框架与动态能力理论，本研究构建以数字化工作活动平台为核心的再造方案：依ESIA原则对复核、签章、会签与清标等环节实施清除、简化、整合与自动化；设计并开发以“项目—时间轴—里程碑—环节—待办”五级流程对象为主干的平台，通过文档解析、清标辅助、风险识别与合规检查等模块承载全周期追踪，并以私有化部署与关键操作留痕满足数据安全与可追溯要求。",
"效果评估从流程效率、集成水平、资源节约、业务价值与组织能力五维度构建指标体系，采用再造前后数据对比与情景测算评估。结果表明，再造后流程周期与清标耗时显著下降，集成与组织能力相关指标达到较高运行水平，环节再造与平台建设对效率形成支撑。受数据治理约束，部分指标缺少再造前基线，未开展统计检验，结论以同期变化与方向性判断为主。",
"本研究将业务流程再造理论落实到可运行的系统与可测量的数据结构，为同类机构在合规约束下推进流程再造与效果评估提供参考。",
]
KEYWORDS = "关键词：工作流程数字化再造；招标代理；流程优化；业务流程再造"

cn = sum(len(re.findall(r"[\u4e00-\u9fa5]", p)) for p in PARAS)
print("摘要汉字数:", cn, "| 段数:", len(PARAS))

d = Document()
sec = d.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.left_margin = sec.right_margin = Cm(3)
sec.top_margin = sec.bottom_margin = Cm(3)
sec.header_distance = Cm(2); sec.footer_distance = Cm(2)

st = d.styles["Normal"]
st.font.name = TNR; st.font.size = Pt(12)
st.element.rPr.rFonts.set(qn("w:eastAsia"), SONG)


def set_run(r, cjk, latin, size, bold=False):
    r.font.name = latin; r.font.size = Pt(size); r.bold = bold
    rPr = r._element.get_or_add_rPr()
    rf = rPr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts"); rPr.append(rf)
    rf.set(qn("w:ascii"), latin); rf.set(qn("w:hAnsi"), latin); rf.set(qn("w:eastAsia"), cjk)


# 摘要标题（标题1：黑体小三居中，1.5倍行距，段后11磅）
h = d.add_paragraph()
h.alignment = WD_ALIGN_PARAGRAPH.CENTER
h.paragraph_format.first_line_indent = Pt(0)
h.paragraph_format.line_spacing = 1.5
h.paragraph_format.space_before = Pt(0)
h.paragraph_format.space_after = Pt(11)
set_run(h.add_run("摘    要"), HEI, ARIAL, 15, bold=True)

# 正文（宋体/TNR 小四，1.25 倍行距，首行缩进 2 字）
for t in PARAS:
    p = d.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf = p.paragraph_format
    pf.first_line_indent = Pt(24)
    pf.line_spacing = 1.25
    pf.space_before = Pt(0); pf.space_after = Pt(0)
    set_run(p.add_run(t), SONG, TNR, 12)

# 关键词（黑体小四加粗，与摘要空一行）
p = d.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.LEFT
pf = p.paragraph_format
pf.first_line_indent = Pt(24)
pf.line_spacing = 1.25
pf.space_before = Pt(12); pf.space_after = Pt(0)
set_run(p.add_run(KEYWORDS), HEI, ARIAL, 12, bold=True)

# 页脚页码（宋体五号居中）
fp = sec.footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
fp.paragraph_format.first_line_indent = Pt(0)
fld = OxmlElement("w:fldSimple"); fld.set(qn("w:instr"), "PAGE")
r = OxmlElement("w:r"); rPr = OxmlElement("w:rPr")
rf = OxmlElement("w:rFonts")
rf.set(qn("w:ascii"), TNR); rf.set(qn("w:hAnsi"), TNR); rf.set(qn("w:eastAsia"), SONG)
sz = OxmlElement("w:sz"); sz.set(qn("w:val"), "21")
rPr.append(rf); rPr.append(sz); r.append(rPr)
tt = OxmlElement("w:t"); tt.text = "1"; r.append(tt); fld.append(r)
fp._p.append(fld)

out = "v12/摘要_700字_新版.docx"
d.save(out)
print("saved", out)
