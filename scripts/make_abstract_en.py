# -*- coding: utf-8 -*-
"""生成与 712 字中文摘要对应的英文摘要（单独 docx，按《格式.doc》英文摘要格式）。"""
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

TNR = "Times New Roman"
SONG = "宋体"

PARAS = [
"Against the backdrop of the digital economy and the deepening full-process electronization of tendering and bidding, the digital re-engineering of workflows has become an important path for tendering agencies to improve efficiency and compliance. Tendering agency business is document-intensive and subject to high compliance requirements; traditional manual review, paper-based circulation and serial countersigning can hardly adapt to the rapid iteration of regulatory standards and clients' rising time requirements. This study takes the tendering agency business line of Company Z as its object and examines the theoretical framework, implementation path and effect evaluation of its workflow digital re-engineering.",
"The study adopts literature review, case study, questionnaire survey, semi-structured interviews, design science research and scenario analysis, and triangulates questionnaires, interviews, institutional documents, system design documents and company-approved aggregate data to form a mixed-method analytical framework.",
"The study finds that Company Z's workflows have been preliminarily digitalized, but three prominent problems remain: digital tools are poorly interconnected, with data silos and duplicated data entry; process nodes are redundant and schedule control is weak, so review and paper-based circulation consume considerable non-value-added time; and employees' digital awareness and operational skills are insufficient, which limits actual adoption of new tools. These problems stem from the interaction of lagging internal management mechanisms, weak technical infrastructure, process design misaligned with external standards, and an inadequate training system, which reinforce one another.",
"Drawing on Business Process Re-engineering theory, the TOE framework and dynamic capability theory, the study develops a re-engineering scheme centred on a digital work-activity platform. Following ESIA principles, redundant steps in review, signing, countersigning and bid clearing are eliminated, simplified, integrated and automated; a platform is then designed and developed around a five-level process object model of \u201cproject\u2013timeline\u2013milestone\u2013step\u2013to-do\u201d, carrying full-cycle tracking through modules for document parsing, bid-clearing assistance, risk identification and compliance checking, while private deployment and auditing of key operations ensure data security and process traceability.",
"The effect evaluation builds an indicator system across five dimensions, namely process efficiency, integration level, resource saving, business value and organizational capability, and combines before-and-after aggregate data comparison with scenario estimation. The results show that process cycle and bid-clearing time decreased significantly after re-engineering, and indicators of integration and organizational capability reached a relatively high operating level, indicating that step-level re-engineering and platform construction positively support operational efficiency. Constrained by corporate data governance, some indicators lack pre-re-engineering baselines and no statistical tests were conducted; the conclusions therefore rest on contemporaneous changes and directional judgements.",
"This study translates Business Process Re-engineering theory into an operational system design and measurable data structures, providing a reference for similar organizations to advance process re-engineering and effect evaluation under compliance constraints.",
]
KW = "Key Words: Workflow Digital Re-engineering; Tendering Agency; Process Optimization; Business Process Re-engineering"

words = sum(len(p.split()) for p in PARAS)
print("正文词数:", words)

d = Document()
sec = d.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.left_margin = sec.right_margin = Cm(3)
sec.top_margin = sec.bottom_margin = Cm(3)
sec.header_distance = Cm(2); sec.footer_distance = Cm(2)

st = d.styles["Normal"]
st.font.name = TNR; st.font.size = Pt(12)
st.element.rPr.rFonts.set(qn("w:eastAsia"), SONG)


def set_run(r, size, bold=False, cjk=TNR):
    r.font.name = TNR; r.font.size = Pt(size); r.bold = bold
    rPr = r._element.get_or_add_rPr()
    rf = rPr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts"); rPr.append(rf)
    rf.set(qn("w:ascii"), TNR); rf.set(qn("w:hAnsi"), TNR); rf.set(qn("w:eastAsia"), cjk)


# 标题 ABSTRACT：TNR 小三 15pt、加粗、居中、1.5 倍行距、段后 11 磅
h = d.add_paragraph()
h.alignment = WD_ALIGN_PARAGRAPH.CENTER
h.paragraph_format.first_line_indent = Pt(0)
h.paragraph_format.line_spacing = 1.5
h.paragraph_format.space_before = Pt(0)
h.paragraph_format.space_after = Pt(11)
set_run(h.add_run("ABSTRACT"), 15, bold=True)

# 正文：TNR 小四 12pt、1.25 倍行距、首行缩进 2 字符（24pt）、两端对齐
for t in PARAS:
    p = d.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf = p.paragraph_format
    pf.first_line_indent = Pt(24)
    pf.line_spacing = 1.25
    pf.space_before = Pt(0); pf.space_after = Pt(0)
    set_run(p.add_run(t), 12)

# Key Words：TNR 小四，标签加粗，与正文空一行
p = d.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
pf = p.paragraph_format
pf.first_line_indent = Pt(0)
pf.line_spacing = 1.25
pf.space_before = Pt(12); pf.space_after = Pt(0)
label, rest = KW.split(":", 1)
set_run(p.add_run(label + ":"), 12, bold=True)
set_run(p.add_run(rest), 12)

# 页脚页码
fp = sec.footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
fp.paragraph_format.first_line_indent = Pt(0)
fld = OxmlElement("w:fldSimple"); fld.set(qn("w:instr"), "PAGE")
r = OxmlElement("w:r"); rPr = OxmlElement("w:rPr")
rf = OxmlElement("w:rFonts"); rf.set(qn("w:ascii"), TNR); rf.set(qn("w:hAnsi"), TNR); rf.set(qn("w:eastAsia"), SONG)
sz = OxmlElement("w:sz"); sz.set(qn("w:val"), "21")
rPr.append(rf); rPr.append(sz); r.append(rPr)
t = OxmlElement("w:t"); t.text = "1"; r.append(t); fld.append(r)
fp._p.append(fld)

out = "v12/ABSTRACT_英文摘要_新版.docx"
d.save(out)
print("saved", out)
